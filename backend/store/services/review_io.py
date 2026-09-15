"""Reading and writing the review book as a spreadsheet.

The catalogue's 1,500-odd reviews arrived as a Judge.me export and the client
still maintains them that way, so the admin screen needs both directions: a file
that opens in Excel, and an importer that swallows a file back — either the one
we just handed out (edited in place, matched on ``id``) or a fresh Judge.me
export (matched on handle + reviewer + body, exactly like the CLI importer).

``import_shopify_reviews`` predates this module and still owns the CLI path; the
handle aliases live here so both read the same map.
"""
import csv
import datetime
import io
import re

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime

from ..models import Product, Review
from .reviews import recalculate_product_review_aggregates


# Some products were re-slugged after their reviews were first imported, so the
# handle in a Shopify export no longer resolves to a Product. Verified by
# reviewer-name overlap against the rows already in the database.
HANDLE_ALIASES = {
    "enfant-organic-plus-moisture-conditioner-for-kids": "Enfant-Organic-Kids-Hair-Conditioner",
    "best-newborn-gift-set-uae-relaxing-night-routine": "newborn-baby-gift-set",
    "enfant-organic-plus-extra-mild-face-body-wipes": "Enfant-Organic-Plus-Extra-Mild-Wipes",
    "enfant-organic-body-wash-shampoo-500-ml": "enfant-organic-body-wash-shampoo",
    "enfant-ultimate-newborn-essential-kit-uae-and-oman": "organic-newborn-essential-kit",
    "enfant-ultra-care-organic-plus-shampoo-body-wash-uae-oman": "Ultra-Care-Shampoo",
}

# ``id`` leads so that a file the admin exported, edited and re-imported updates
# those same rows instead of duplicating them.
REVIEW_EXPORT_HEADERS = (
    "id",
    "product_slug",
    "product_name",
    "customer_name",
    "rating",
    "title",
    "comment",
    "images",
    "is_approved",
    "is_verified_purchase",
    "order_number",
    "created_at",
)

# A spreadsheet from the client will not use our column names, and Excel mangles
# case and spacing on the way through. Headers are normalised to lowercase words
# joined by underscores before being looked up here.
COLUMN_ALIASES = {
    "id": "id",
    "review_id": "id",
    "product_slug": "product",
    "product_handle": "product",
    "handle": "product",
    "product": "product",
    "product_id": "product",
    "product_name": "product_name",
    "customer_name": "customer_name",
    "reviewer_name": "customer_name",
    "author": "customer_name",
    "name": "customer_name",
    "rating": "rating",
    "score": "rating",
    "stars": "rating",
    "title": "title",
    "headline": "title",
    "review_title": "title",
    "comment": "comment",
    "body": "comment",
    "review": "comment",
    "content": "comment",
    "review_body": "comment",
    "images": "images",
    "media_urls": "images",
    "photos": "images",
    "picture_urls": "images",
    "is_approved": "is_approved",
    "approved": "is_approved",
    "published": "is_approved",
    "status": "is_approved",
    "is_verified_purchase": "is_verified_purchase",
    "verified": "is_verified_purchase",
    "verified_purchase": "is_verified_purchase",
    "order_number": "order_number",
    "created_at": "created_at",
    "date": "created_at",
    "review_date": "created_at",
    "submitted_at": "created_at",
}

TRUE_WORDS = {"1", "true", "yes", "y", "on", "approved", "published", "publish", "live", "active"}
FALSE_WORDS = {"0", "false", "no", "n", "off", "pending", "unpublished", "hidden", "spam", "removed", "rejected"}

# An import is a paste of a spreadsheet, not a data feed: these ceilings exist so
# a wrong file (a 200MB image dump, a runaway generated CSV) is refused in one
# cheap check instead of holding a gunicorn worker for minutes.
MAX_IMPORT_BYTES = 12 * 1024 * 1024
MAX_IMPORT_ROWS = 20000

# Errors are reported back to the admin screen; past a handful they stop being
# read and start being scrolled.
MAX_REPORTED_ERRORS = 25


class ReviewImportError(Exception):
    """The file cannot be read at all — wrong type, empty, or far too big."""


# ─── Export ───────────────────────────────────────────────────────────────────


def review_export_queryset(queryset=None):
    """Oldest first, with everything the row builder touches already joined."""
    queryset = queryset if queryset is not None else Review.objects.all()
    return queryset.select_related("product", "order").order_by("created_at", "id")


def _review_export_row(review):
    created = review.created_at
    return [
        review.id,
        review.product.slug if review.product_id else "",
        review.product.name_en if review.product_id else "",
        review.customer_name,
        review.rating,
        review.title,
        review.comment,
        # One cell, one URL per line: Excel keeps the newlines and the importer
        # splits on them again, so a round-trip does not lose photos.
        "\n".join(review.images if isinstance(review.images, list) else []),
        "yes" if review.is_approved else "no",
        "yes" if review.is_verified_purchase else "no",
        review.order.order_number if review.order_id else "",
        timezone.localtime(created).isoformat(timespec="seconds") if created else "",
    ]


def review_export_rows(queryset=None):
    for review in review_export_queryset(queryset).iterator(chunk_size=500):
        yield _review_export_row(review)


def review_export_filename(*, prefix="enfant_reviews"):
    return f"{prefix}_{timezone.localdate().isoformat()}"


def review_export_response(queryset=None, *, export_format="csv", filename=None):
    """The review book as a CSV or an Excel workbook."""
    from django.http import HttpResponse

    filename = filename or review_export_filename()
    headers = list(REVIEW_EXPORT_HEADERS)

    if str(export_format).lower() in {"xlsx", "excel"}:
        import openpyxl
        from openpyxl.utils import get_column_letter

        workbook = openpyxl.Workbook(write_only=True)
        sheet = workbook.create_sheet("Reviews")
        for index, header in enumerate(headers, start=1):
            # The comment column carries whole paragraphs; everything else reads
            # better narrow.
            width = 60 if header == "comment" else max(14, len(header) + 4)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.append(headers)
        for row in review_export_rows(queryset):
            sheet.append(row)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in review_export_rows(queryset):
        writer.writerow(row)
    return response


# ─── Import ───────────────────────────────────────────────────────────────────


def _normalise_header(value):
    text = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower())
    return text.strip("_")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_bool(value, default=None):
    text = _cell_text(value).lower()
    if not text:
        return default
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    return default


def _as_rating(value):
    text = _cell_text(value)
    if not text:
        return None
    try:
        rating = int(round(float(text)))
    except (TypeError, ValueError):
        return None
    return max(1, min(5, rating))


def _as_images(value):
    if isinstance(value, (list, tuple)):
        candidates = [str(item) for item in value]
    else:
        candidates = re.split(r"[\n,;|]+", _cell_text(value))
    images = []
    for candidate in candidates:
        url = candidate.strip().strip('"')
        # Admin-uploaded photos are stored as /media/… paths, so a round-trip of
        # our own export has to survive alongside Judge.me's absolute URLs.
        if url.startswith(("http://", "https://", "/")) and url not in images:
            images.append(url)
    return images


def _as_datetime(value):
    if isinstance(value, datetime.datetime):
        parsed = value
    elif isinstance(value, datetime.date):
        parsed = datetime.datetime(value.year, value.month, value.day)
    else:
        text = _cell_text(value)
        if not text:
            return None
        parsed = parse_datetime(text)
        if parsed is None:
            as_date = parse_date(text[:10])
            if as_date is None:
                return None
            parsed = datetime.datetime(as_date.year, as_date.month, as_date.day)
    if timezone.is_naive(parsed):
        return timezone.make_aware(parsed)
    return parsed


def _read_table(file_obj, filename=""):
    """Return ``(headers, rows)`` from a .csv or .xlsx upload."""
    raw = file_obj.read()
    if isinstance(raw, str):
        raw = raw.encode("utf-8")
    if not raw:
        raise ReviewImportError("The file is empty.")
    if len(raw) > MAX_IMPORT_BYTES:
        raise ReviewImportError(
            f"The file is {len(raw) // (1024 * 1024)}MB. Reviews import accepts files up to "
            f"{MAX_IMPORT_BYTES // (1024 * 1024)}MB."
        )

    name = str(filename or getattr(file_obj, "name", "") or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises a zoo of exception types
            raise ReviewImportError(f"That .xlsx could not be opened: {exc}") from exc
        sheet = workbook.active
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith(".csv") or not name:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            # Excel on Windows still writes cp1252 for anything non-ASCII.
            text = raw.decode("cp1252", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        table = [row for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ReviewImportError("Unsupported file type. Upload a .csv or .xlsx file.")

    table = [row for row in table if any(_cell_text(cell) for cell in row)]
    if not table:
        raise ReviewImportError("The file has no rows.")
    if len(table) - 1 > MAX_IMPORT_ROWS:
        raise ReviewImportError(
            f"The file has {len(table) - 1} rows. Reviews import accepts up to {MAX_IMPORT_ROWS}."
        )

    headers = [_normalise_header(cell) for cell in table[0]]
    if not any(headers):
        raise ReviewImportError("The first row must be a header row.")
    return headers, table[1:]


def _row_to_fields(headers, row):
    """Map one spreadsheet row onto canonical field names."""
    fields = {}
    for index, header in enumerate(headers):
        field = COLUMN_ALIASES.get(header)
        if not field:
            continue
        value = row[index] if index < len(row) else None
        # First column wins: a Judge.me export carries both "Body" and "Review
        # Body", and the second is usually the empty one.
        if field not in fields or (fields[field] in (None, "") and value not in (None, "")):
            fields[field] = value
    return fields


def _product_resolver():
    """Handle/id/name → Product pk, tolerating case and post-import re-slugs."""
    slug_to_pk = dict(Product.objects.values_list("slug", "pk"))
    lower_slug_to_pk = {slug.lower(): pk for slug, pk in slug_to_pk.items()}
    pks = set(slug_to_pk.values())

    # Names are only used when no handle resolves, and only when the name picks
    # out exactly one product — two products sharing a name would otherwise send
    # reviews to whichever row the database happened to return first.
    name_counts = {}
    for pk, name in Product.objects.values_list("pk", "name_en"):
        key = (name or "").strip().lower()
        if not key:
            continue
        name_counts.setdefault(key, []).append(pk)
    name_to_pk = {key: values[0] for key, values in name_counts.items() if len(values) == 1}

    def resolve(handle, name=""):
        text = _cell_text(handle)
        if text:
            for candidate in (text, HANDLE_ALIASES.get(text, "")):
                if not candidate:
                    continue
                pk = slug_to_pk.get(candidate) or lower_slug_to_pk.get(candidate.lower())
                if pk:
                    return pk
            if text.isdigit() and int(text) in pks:
                return int(text)
        name_key = _cell_text(name).lower()
        if name_key:
            return name_to_pk.get(name_key)
        return None

    return resolve


def dedup_key_for(product_pk, customer_name, title, comment):
    """Identity of a review in a file that carries no ids.

    Reviewer names arrive anonymised ("r***a") and titles are often blank, so
    name+title alone collapses genuinely different reviews of the same product.
    The body disambiguates them. Kept byte-compatible with the CLI importer so
    the two paths agree on what counts as already-imported.
    """
    return (
        product_pk,
        (customer_name or "").lower(),
        (title or "")[:160].lower(),
        (comment or "")[:200].strip().lower(),
    )


def import_reviews(
    file_obj,
    *,
    filename="",
    dry_run=False,
    default_approved=True,
    update_existing=True,
    backfill_images=True,
):
    """Load a review spreadsheet into the Review table.

    ``default_approved`` only decides rows whose file says nothing about status;
    a Status/approved column in the file always wins for its own row.

    ``update_existing`` covers the round-trip: a row carrying an ``id`` we know
    edits that review in place. Rows without an id are matched on content and
    skipped if already present, so re-running a Judge.me export is safe —
    that is how the catalogue picks up reviews for products added since the last
    run.
    """
    headers, rows = _read_table(file_obj, filename)

    known = {COLUMN_ALIASES.get(header) for header in headers}
    if "product" not in known and "id" not in known:
        raise ReviewImportError(
            "No product column found. The file needs a 'product_slug' (or Judge.me "
            "'Product Handle') column, or an 'id' column to update existing reviews."
        )

    resolve_product = _product_resolver()

    existing_keys = {}
    for review in Review.objects.values("pk", "product_id", "customer_name", "title", "comment"):
        existing_keys.setdefault(
            dedup_key_for(
                review["product_id"], review["customer_name"], review["title"], review["comment"]
            ),
            review["pk"],
        )
    existing_ids = set(Review.objects.values_list("pk", flat=True))

    stats = {
        "rows": 0,
        "created": 0,
        "updated": 0,
        "skipped_duplicate": 0,
        "skipped_no_product": 0,
        "skipped_invalid": 0,
        "images_backfilled": 0,
        "errors": 0,
    }
    messages = []
    unmatched_products = {}
    affected_pks = set()
    created_dates = []  # (pk, created_at) — auto_now_add blocks setting it at create time

    def note(row_number, message):
        stats["errors"] += 1
        if len(messages) < MAX_REPORTED_ERRORS:
            messages.append(f"Row {row_number}: {message}")

    with transaction.atomic():
        for offset, raw_row in enumerate(rows):
            row_number = offset + 2  # +1 for the header, +1 for 1-based rows
            stats["rows"] += 1
            fields = _row_to_fields(headers, raw_row)

            review_id = _cell_text(fields.get("id"))
            rating = _as_rating(fields.get("rating"))
            customer_name = _cell_text(fields.get("customer_name"))[:160]
            title = _cell_text(fields.get("title"))[:160]
            comment = _cell_text(fields.get("comment"))
            images = _as_images(fields.get("images"))
            approved = _as_bool(fields.get("is_approved"), None)
            verified = _as_bool(fields.get("is_verified_purchase"), False)
            created_at = _as_datetime(fields.get("created_at"))

            # ── Update an existing review, matched on the id our export wrote ──
            if review_id.isdigit() and int(review_id) in existing_ids:
                if not update_existing:
                    stats["skipped_duplicate"] += 1
                    continue
                update = {}
                if customer_name:
                    update["customer_name"] = customer_name
                if rating is not None:
                    update["rating"] = rating
                if "title" in fields:
                    update["title"] = title
                if comment:
                    update["comment"] = comment
                if images:
                    update["images"] = images
                if approved is not None:
                    update["is_approved"] = approved
                if fields.get("is_verified_purchase") not in (None, ""):
                    update["is_verified_purchase"] = bool(verified)
                product_pk = resolve_product(fields.get("product"), fields.get("product_name"))
                if product_pk:
                    update["product_id"] = product_pk
                if not update:
                    stats["skipped_duplicate"] += 1
                    continue
                review = Review.objects.filter(pk=int(review_id)).first()
                if review is None:  # deleted between the id scan and here
                    stats["skipped_invalid"] += 1
                    continue
                affected_pks.add(review.product_id)
                for field, value in update.items():
                    setattr(review, field, value)
                review.save(update_fields=list(update.keys()) + ["updated_at"])
                if created_at:
                    Review.objects.filter(pk=review.pk).update(created_at=created_at)
                affected_pks.add(review.product_id)
                stats["updated"] += 1
                continue

            if review_id and not review_id.isdigit():
                note(row_number, f"'{review_id}' is not a review id.")
                stats["skipped_invalid"] += 1
                continue

            # ── Otherwise it is a new review, keyed on its content ────────────
            product_pk = resolve_product(fields.get("product"), fields.get("product_name"))
            if product_pk is None:
                handle = _cell_text(fields.get("product")) or _cell_text(fields.get("product_name")) or "(blank)"
                unmatched_products[handle] = unmatched_products.get(handle, 0) + 1
                stats["skipped_no_product"] += 1
                continue

            if rating is None:
                note(row_number, "rating is missing or not a number 1-5.")
                stats["skipped_invalid"] += 1
                continue
            if not comment and not title:
                note(row_number, "the review has neither a title nor a comment.")
                stats["skipped_invalid"] += 1
                continue

            customer_name = customer_name or "Anonymous"
            key = dedup_key_for(product_pk, customer_name, title, comment)
            if key in existing_keys:
                stats["skipped_duplicate"] += 1
                if backfill_images and images:
                    # Reviews imported before photo support have none; give them
                    # the file's media without touching their text.
                    stats["images_backfilled"] += Review.objects.filter(
                        pk=existing_keys[key], images=[]
                    ).update(images=images)
                continue

            try:
                review = Review.objects.create(
                    product_id=product_pk,
                    customer_name=customer_name,
                    rating=rating,
                    title=title,
                    comment=comment,
                    images=images,
                    is_approved=default_approved if approved is None else approved,
                    is_verified_purchase=bool(verified),
                )
            except Exception as exc:
                note(row_number, str(exc))
                stats["skipped_invalid"] += 1
                continue

            existing_keys[key] = review.pk
            existing_ids.add(review.pk)
            affected_pks.add(product_pk)
            stats["created"] += 1
            if created_at:
                created_dates.append((review.pk, created_at))

        for pk, created_at in created_dates:
            Review.objects.filter(pk=pk).update(created_at=created_at)

        for product_pk in affected_pks:
            recalculate_product_review_aggregates(product_pk)

        if dry_run:
            transaction.set_rollback(True)

    stats["dry_run"] = bool(dry_run)
    stats["products_touched"] = len(affected_pks)
    stats["messages"] = messages
    # A handle no product answers to means those reviews are silently lost —
    # surface them instead of hiding them inside a count.
    stats["unmatched_products"] = [
        {"value": handle, "rows": count}
        for handle, count in sorted(unmatched_products.items(), key=lambda item: -item[1])[:20]
    ]
    return stats
